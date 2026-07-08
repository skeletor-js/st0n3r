"""Document-format tests: EPUB3 (U3, stdlib, always available), PDF (U4,
reportlab behind the export extra), and Shunn DOCX (U5, python-docx behind
the export extra). PDF/DOCX tests importorskip their dependency; a
missing-extra monkeypatch test proves each degrades independently while EPUB
still builds. The novella corpus is read-only (copied to tmp_path)."""

from __future__ import annotations

import shutil
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest

from stoner.canon.scaffold import scaffold_project
from stoner.ledger import Ledger
from stoner.project import WritingProject

CORPUS = Path(__file__).resolve().parent.parent / "examples" / "novella"


def _make_project(tmp_path: Path, chapters: int = 3) -> WritingProject:
    p = WritingProject.create(tmp_path / "book", "book")
    scaffold_project(p, "book")
    p.config.ship.title = "Test Book"
    p.config.ship.author = "A. Writer"
    for n in range(1, chapters + 1):
        p.write_chapter(
            n,
            {"title": f"Chapter {n} Title", "status": "revised"},
            f"First paragraph of chapter {n}.\n\nSecond one with *emphasis* & <angle>.",
        )
    return p


def corpus_copy(tmp_path: Path) -> WritingProject:
    dest = tmp_path / "novella"
    shutil.copytree(CORPUS, dest)
    return WritingProject(dest)


# ---------------------------------------------------------------------------
# U3: EPUB3
# ---------------------------------------------------------------------------


def test_epub_structure(tmp_path: Path):
    from stoner.ship.epub import write_epub

    project = _make_project(tmp_path, chapters=3)
    path, chapters = write_epub(project)
    assert chapters == 3
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        # mimetype must be the first entry, stored uncompressed, exact content.
        assert names[0] == "mimetype"
        info = zf.getinfo("mimetype")
        assert info.compress_type == zipfile.ZIP_STORED
        assert zf.read("mimetype") == b"application/epub+zip"

        container = ET.fromstring(zf.read("META-INF/container.xml"))
        rootfile = container.find(".//{*}rootfile")
        assert rootfile is not None
        assert rootfile.attrib["full-path"] == "OEBPS/package.opf"

        opf = ET.fromstring(zf.read("OEBPS/package.opf"))
        spine = opf.find("{*}spine")
        assert spine is not None
        chapter_refs = [i for i in spine if i.tag.endswith("itemref")]
        # 3 front-matter pages + 3 chapters + 1 end page all spine-ordered.
        assert len(chapter_refs) == len(project.chapters()) + 4

        nav = zf.read("OEBPS/nav.xhtml").decode("utf-8")
        assert "Chapter 1 Title" in nav
        assert "Chapter 3 Title" in nav


def test_epub_byte_identical_rebuild(tmp_path: Path):
    from stoner.ship.epub import write_epub

    project = _make_project(tmp_path)
    p1, _ = write_epub(project)
    first = p1.read_bytes()
    p2, _ = write_epub(project)
    assert p2.read_bytes() == first


def test_epub_escaping_and_italics(tmp_path: Path):
    from stoner.ship.epub import write_epub

    project = _make_project(tmp_path, chapters=1)
    path, _ = write_epub(project)
    with zipfile.ZipFile(path) as zf:
        chap = zf.read("OEBPS/chap-001.xhtml").decode("utf-8")
    assert "<em>emphasis</em>" in chap
    assert "&amp;" in chap and "&lt;angle&gt;" in chap
    assert "<angle>" not in chap


def test_epub_corpus_integration(tmp_path: Path):
    from stoner.ship.epub import write_epub

    project = corpus_copy(tmp_path)
    path, chapters = write_epub(project)
    assert chapters == 15
    with zipfile.ZipFile(path) as zf:
        assert zipfile.ZipFile.testzip(zf) is None
    tail = Ledger(project.root).tail(1)[0]
    assert tail.action == "ship.epub"


def test_epub_identifier_is_deterministic():
    from stoner.ship.epub import book_identifier

    a = book_identifier("novella", "Sungrown")
    b = book_identifier("novella", "Sungrown")
    c = book_identifier("novella", "Other")
    assert a == b and a != c
    assert a.startswith("urn:uuid:")


# ---------------------------------------------------------------------------
# U4: PDF (reportlab, export extra)
# ---------------------------------------------------------------------------


def test_pdf_trim_parsing():
    from stoner.ship.pdf import _parse_trim

    assert _parse_trim("5.5x8.5") == (396.0, 612.0)
    assert _parse_trim("6x9") == (432.0, 648.0)
    assert _parse_trim("garbage") == (396.0, 612.0)


def _pdf_page_count(data: bytes) -> int:
    import re

    return len(re.findall(rb"/Type\s*/Page(?![s])", data))


def test_pdf_builds_and_paginates(tmp_path: Path):
    pytest.importorskip("reportlab")
    from stoner.ship.pdf import write_pdf

    project = _make_project(tmp_path, chapters=3)
    path, chapters = write_pdf(project)
    assert chapters == 3
    data = path.read_bytes()
    assert data[:4] == b"%PDF"
    assert _pdf_page_count(data) > 0
    assert b"/MediaBox [ 0 0 396 612 ]" in data


def test_pdf_page_count_grows_with_chapters(tmp_path: Path):
    pytest.importorskip("reportlab")
    from stoner.ship.pdf import write_pdf

    small = _make_project(tmp_path / "a", chapters=1)
    big = _make_project(tmp_path / "b", chapters=6)
    ps = _pdf_page_count(write_pdf(small)[0].read_bytes())
    pb = _pdf_page_count(write_pdf(big)[0].read_bytes())
    assert pb > ps


def test_pdf_missing_extra_degrades_but_epub_works(tmp_path: Path, monkeypatch):
    import builtins

    from stoner.ship.epub import write_epub
    from stoner.ship.manifest import ShipError
    from stoner.ship.pdf import write_pdf

    project = _make_project(tmp_path, chapters=2)
    real_import = builtins.__import__

    def fake_import(name, *args, **kwargs):
        if name.startswith("reportlab"):
            raise ImportError("simulated missing reportlab")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fake_import)
    with pytest.raises(ShipError, match=r"pip install 'st0n3r\[export\]'"):
        write_pdf(project)
    # EPUB in the same project still succeeds (R9).
    monkeypatch.setattr(builtins, "__import__", real_import)
    path, chapters = write_epub(project)
    assert chapters == 2 and path.exists()


def test_pdf_corpus_integration(tmp_path: Path):
    pytest.importorskip("reportlab")
    from stoner.ship.pdf import write_pdf

    project = corpus_copy(tmp_path)
    path, chapters = write_pdf(project)
    assert chapters == 15 and path.read_bytes()[:4] == b"%PDF"
    tail = Ledger(project.root).tail(1)[0]
    assert tail.action == "ship.pdf"


# ---------------------------------------------------------------------------
# U5: Shunn DOCX (python-docx, export extra)
# ---------------------------------------------------------------------------


def _docx_project(tmp_path: Path, chapters: int = 2) -> WritingProject:
    p = _make_project(tmp_path, chapters=chapters)
    p.config.ship.contact_lines = ["Jane Author", "123 Main St", "jane@example.com"]
    return p


def test_docx_margins_font_spacing(tmp_path: Path):
    pytest.importorskip("docx")
    from docx import Document
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Inches, Pt

    from stoner.ship.shunn import write_docx

    project = _docx_project(tmp_path)
    path, chapters = write_docx(project)
    assert chapters == 2
    doc = Document(str(path))
    section = doc.sections[0]
    assert section.left_margin == Inches(1)
    assert section.top_margin == Inches(1)
    normal = doc.styles["Normal"]
    assert normal.font.name == "Times New Roman"
    assert normal.font.size == Pt(12)
    assert normal.paragraph_format.line_spacing_rule == WD_LINE_SPACING.DOUBLE


def test_docx_header_has_page_field(tmp_path: Path):
    pytest.importorskip("docx")
    from stoner.ship.shunn import write_docx

    project = _docx_project(tmp_path)
    project.config.ship.author = "A. Writer"
    project.config.ship.title = "Test Book"
    path, _ = write_docx(project)
    with zipfile.ZipFile(path) as zf:
        header_parts = [n for n in zf.namelist() if n.startswith("word/header")]
        blob = "".join(zf.read(n).decode("utf-8") for n in header_parts)
    assert "instrText" in blob and "PAGE" in blob
    assert "Writer" in blob and "Test Book" in blob


def test_docx_underline_italics_flag(tmp_path: Path):
    pytest.importorskip("docx")
    from docx import Document

    from stoner.ship.shunn import write_docx

    def italic_underline_runs(path):
        doc = Document(str(path))
        italic, underline = False, False
        for p in doc.paragraphs:
            for r in p.runs:
                if r.italic:
                    italic = True
                if r.underline:
                    underline = True
        return italic, underline

    # Default: emphasis is italic.
    project = _docx_project(tmp_path / "a", chapters=1)
    project.write_chapter(1, {"title": "One", "status": "revised"}, "A word *emph* here.")
    path, _ = write_docx(project)
    it, un = italic_underline_runs(path)
    assert it and not un

    # underline_italics=True: emphasis comes back underlined, not italic.
    project2 = _docx_project(tmp_path / "b", chapters=1)
    project2.config.ship.underline_italics = True
    project2.write_chapter(1, {"title": "One", "status": "revised"}, "A word *emph* here.")
    path2, _ = write_docx(project2)
    it2, un2 = italic_underline_runs(path2)
    assert un2 and not it2


def test_docx_word_count_rounded(tmp_path: Path):
    pytest.importorskip("docx")
    from docx import Document

    from stoner.ship.shunn import write_docx

    project = _docx_project(tmp_path, chapters=1)
    # 12 words -> rounds to 0? use a body with a known-ish count near 500.
    body = " ".join(["word"] * 640)
    project.write_chapter(1, {"title": "One", "status": "revised"}, body)
    path, _ = write_docx(project)
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "approximately 500 words" in text  # 640 rounds to nearest 500


def test_docx_missing_extra_degrades(tmp_path: Path, monkeypatch):
    import builtins

    from stoner.ship.epub import write_epub
    from stoner.ship.manifest import ShipError
    from stoner.ship.shunn import write_docx

    project = _docx_project(tmp_path, chapters=2)
    real_import = builtins.__import__

    def fake_import(name, *args, **kwargs):
        if name == "docx" or name.startswith("docx."):
            raise ImportError("simulated missing python-docx")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fake_import)
    with pytest.raises(ShipError, match=r"pip install 'st0n3r\[export\]'"):
        write_docx(project)
    monkeypatch.setattr(builtins, "__import__", real_import)
    _path, chapters = write_epub(project)
    assert chapters == 2


def test_docx_corpus_integration(tmp_path: Path):
    pytest.importorskip("docx")
    from docx import Document

    from stoner.ship.shunn import write_docx

    project = corpus_copy(tmp_path)
    path, chapters = write_docx(project)
    assert chapters == 15
    doc = Document(str(path))
    headings = [p.text for p in doc.paragraphs if p.text in ("Cloning", "The Queue", "Chapter 6")]
    assert "Cloning" in headings
    tail = Ledger(project.root).tail(1)[0]
    assert tail.action == "ship.docx"
