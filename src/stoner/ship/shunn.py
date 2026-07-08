"""Shunn standard-manuscript-format DOCX via python-docx, behind `export` (R8).

Named `shunn.py`, not `docx.py`, to avoid shadowing the library. Same lazy
import + independent degradation contract as the PDF: a missing extra fails
only this format with the `pip install 'st0n3r[export]'` hint (R9).

Modern Shunn is the default (Times New Roman, real italics); the
`ship.underline_italics` flag emits classic Courier-era underlines instead.
python-docx has no first-class page-number field, so the `Surname / TITLE /
page` running header injects a `PAGE` field as raw OOXML field codes
(`fldChar`/`instrText`) -- the documented workaround, isolated in one helper.
`different_first_page` suppresses the header on the title page.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from ..ledger import Ledger
from ..project import WritingProject, count_words
from .assemble import Book, Heading, Paragraph, SceneBreak, assemble_book
from .manifest import ShipError, ShipManifest, require_ready

_INSTALL_HINT = (
    "DOCX export needs the optional python-docx dependency. Install it with:\n"
    "    pip install 'st0n3r[export]'\n"
    "Other formats (epub, pdf) are unaffected."
)


def _round_500(n: int) -> int:
    """Round a word count to the nearest 500 (Shunn's flat approximation)."""
    return int(round(n / 500.0)) * 500


def _surname(author: str) -> str:
    author = author.strip()
    if not author:
        return "Author"
    return author.split()[-1]


def _add_page_field(paragraph: Any) -> None:
    """Append a live `PAGE` field to `paragraph` via raw OOXML field codes.

    python-docx exposes no page-number field (python-openxml/python-docx#686),
    so we drop down to the oxml layer it already carries: a `fldChar` begin, an
    `instrText` of "PAGE", and a `fldChar` end, wrapped in `w:r` runs.
    """
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    begin = run._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "begin"})
    run._r.append(begin)

    instr_run = paragraph.add_run()
    instr = instr_run._r.makeelement(qn("w:instrText"), {qn("xml:space"): "preserve"})
    instr.text = " PAGE "
    instr_run._r.append(instr)

    end_run = paragraph.add_run()
    end = end_run._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "end"})
    end_run._r.append(end)


def build_docx_document(
    book: Book,
    *,
    author: str,
    title: str,
    contact_lines: list[str],
    word_count: int,
    underline_italics: bool = False,
) -> Any:
    """Build and return a python-docx Document in Shunn format. Raises
    ShipError if the export extra is missing."""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
        from docx.shared import Inches, Pt
    except ImportError as e:
        raise ShipError(_INSTALL_HINT) from e

    document = Document()

    # -- page + Normal style ------------------------------------------------
    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Inches(0.5)

    def body_paragraph() -> Any:
        return document.add_paragraph()

    def flat(text: str, *, align: Any = None, indent: bool = False) -> Any:
        p = document.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0.5 if indent else 0)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        if align is not None:
            p.alignment = align
        if text:
            p.add_run(text)
        return p

    # -- running header (page 2+): Surname / TITLE / PAGE -------------------
    section = document.sections[0]
    section.different_first_page_header_footer = True
    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_p.add_run(f"{_surname(author)} / {title} / ")
    _add_page_field(header_p)

    # -- title page ---------------------------------------------------------
    for line in contact_lines:
        flat(line, align=WD_ALIGN_PARAGRAPH.LEFT)
    flat(f"approximately {word_count:,} words", align=WD_ALIGN_PARAGRAPH.RIGHT)
    for _ in range(8):
        flat("")
    flat(title, align=WD_ALIGN_PARAGRAPH.CENTER)
    if author:
        flat(f"by {author}", align=WD_ALIGN_PARAGRAPH.CENTER)

    # -- chapters -----------------------------------------------------------
    for chapter in book.chapters:
        document.add_page_break()
        heading = document.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.paragraph_format.first_line_indent = Inches(0)
        heading.add_run(chapter.title)
        flat("")
        for block in chapter.blocks:
            if isinstance(block, Paragraph):
                p = body_paragraph()
                for run in block.runs:
                    r = p.add_run(run.text)
                    if run.style == "em":
                        if underline_italics:
                            r.underline = True
                        else:
                            r.italic = True
                    elif run.style == "strong":
                        r.bold = True
            elif isinstance(block, SceneBreak):
                flat("#", align=WD_ALIGN_PARAGRAPH.CENTER)
            elif isinstance(block, Heading):
                flat(block.text, align=WD_ALIGN_PARAGRAPH.CENTER)

    flat("END", align=WD_ALIGN_PARAGRAPH.CENTER)
    return document


def write_docx(
    project: WritingProject,
    allow_incomplete: bool = False,
    dest: Path | None = None,
    manifest: ShipManifest | None = None,
) -> tuple[Path, int]:
    """Gate, assemble, and write the Shunn DOCX; ledger `ship.docx`. Returns
    (path, chapter_count). Raises ShipError on blockers or a missing extra.
    A pre-built `manifest` skips the gate (used by `ship all`)."""
    if manifest is None:
        manifest = require_ready(project, allow_incomplete)
    book = assemble_book(project, manifest)
    total_words = 0
    for ref in manifest.chapters:
        _fm, body = project.read_chapter(ref.number)
        total_words += count_words(body)
    document = build_docx_document(
        book,
        author=manifest.author,
        title=manifest.title,
        contact_lines=manifest.contact_lines,
        word_count=_round_500(total_words),
        underline_italics=project.config.ship.underline_italics,
    )
    out = dest or project.resolve(f"export/{manifest.slug}-manuscript.docx")
    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out))
    Ledger(project.root).append(
        "ship.docx", target=str(out.relative_to(project.root)), chapters=len(book.chapters)
    )
    return out, len(book.chapters)
